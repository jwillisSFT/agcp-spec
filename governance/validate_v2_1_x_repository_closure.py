#!/usr/bin/env python3
from __future__ import annotations
import hashlib, json, re, sys, urllib.parse
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / 'governance/AGCP-v2.1.x-repository-closure-manifest.json'
REPORT = ROOT / 'governance/AGCP-v2.1.x-repository-closure-validation.json'
RETIRED_NS = {'NS-8.6A-01','NS-8.6A-03','NS-9.1-01'}

def sha256(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()

def add(checks, name, passed, detail=None):
    checks.append({'check':name,'status':'PASS' if passed else 'FAIL','detail':detail})
    return passed

def j(rel):
    return json.loads((ROOT/rel).read_text(encoding='utf-8'))

def main():
    checks=[]; errors=[]; warnings=[]
    required=[
      'README.md','PACKAGE_README.md','RELEASE_NOTES_v2.1.x.md',
      'governance/AGCP-v2.1.x-REPOSITORY-CLOSURE.md',
      'governance/AGCP-v2.1.x-repository-closure-manifest.json',
      'spec/AGCP_Requirements_Traceability_Matrix_(RTM).xlsx',
      'spec/AGCP-v2.1.x-rtm-synchronization-validation.json',
      'conformance/AGCP-v2.1.x-aggregate-conformance-layer-validation.json',
      'conformance/test-mapping.json','conformance/test-control-mapping.json',
      'conformance/harness-checks.json','conformance/fixture-mapping.json'
    ]
    missing=[x for x in required if not (ROOT/x).is_file()]
    add(checks,'required_closure_artifacts_present',not missing,{'missing':missing})
    errors += [f'missing:{x}' for x in missing]

    version=(ROOT/'VERSION').read_text().strip() if (ROOT/'VERSION').is_file() else None
    notes=(ROOT/'RELEASE_NOTES_v2.1.x.md').read_text(encoding='utf-8') if (ROOT/'RELEASE_NOTES_v2.1.x.md').is_file() else ''
    version_ok=(version=='2.1.0' and 'Publication version:** Not assigned by this closure step' in notes)
    add(checks,'published_baseline_preserved_and_new_publication_version_not_silently_assigned',version_ok,{'VERSION':version})
    if not version_ok: errors.append('release-version-boundary')

    rtm=j('spec/AGCP-v2.1.x-rtm-synchronization-validation.json')
    rtm_ok=(rtm.get('status')=='PASS' and rtm.get('rtm_dataset_version')=='RTM-1.47' and rtm.get('main_rtm_row_count')==122 and rtm.get('active_normative_statement_count')==390 and rtm.get('retired_normative_statement_count')==3 and rtm.get('harness_check_count')==21 and rtm.get('harness_test_vector_count')==71 and not rtm.get('errors') and not rtm.get('warnings'))
    add(checks,'final_rtm_synchronization_passes',rtm_ok,{k:rtm.get(k) for k in ['rtm_dataset_version','main_rtm_row_count','ns_cr_relationship_count','active_normative_statement_count','retired_normative_statement_count','formal_test_case_count','harness_check_count','harness_test_vector_count']})
    if not rtm_ok: errors.append('rtm-validation')

    agg=j('conformance/AGCP-v2.1.x-aggregate-conformance-layer-validation.json')
    agg_ok=(agg.get('status')=='PASS' and agg.get('formal_test_case_count')==122 and agg.get('test_mapping_count')==122 and agg.get('test_control_mapping_count')==122 and agg.get('harness_check_count')==21 and agg.get('harness_test_vector_count')==71 and agg.get('controlled_fixture_count')==36 and not agg.get('errors') and not agg.get('warnings'))
    add(checks,'aggregate_conformance_layer_validation_passes',agg_ok,{k:agg.get(k) for k in ['formal_test_case_count','test_mapping_count','test_control_mapping_count','harness_check_count','harness_test_vector_count','controlled_fixture_count']})
    if not agg_ok: errors.append('aggregate-conformance-validation')

    # Workbook sanity: 122 data rows, RTM-1.47 throughout.
    wb=load_workbook(ROOT/'spec/AGCP_Requirements_Traceability_Matrix_(RTM).xlsx',read_only=True,data_only=False)
    ws=wb[wb.sheetnames[0]]
    headers={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
    dscol=headers.get('Dataset_Version'); crcol=headers.get('CR_ID'); tccol=headers.get('TC_ID')
    rows=[]
    for r in range(2,ws.max_row+1):
        cr=ws.cell(r,crcol).value if crcol else None
        if cr: rows.append(r)
    workbook_ok=(len(rows)==122 and dscol is not None and all(ws.cell(r,dscol).value=='RTM-1.47' for r in rows) and tccol is not None and len({ws.cell(r,tccol).value for r in rows})==122)
    add(checks,'authoritative_rtm_workbook_matches_final_dataset',workbook_ok,{'data_rows':len(rows),'sheet':ws.title,'max_columns':ws.max_column})
    if not workbook_ok: errors.append('rtm-workbook')

    # Current validation summaries must no longer report pending RTM synchronization.
    current_reports=[ROOT/'DS_INTERFACE_API_VALIDATION_SUMMARY.json', ROOT/'api/interface-traceability-validation.json', ROOT/'registries/registry-entry-traceability-validation.json', ROOT/'schemas/catalog/schema-catalog-validation.json']
    current_reports += sorted((ROOT/'schemas').glob('DS-*-v2.1.x-validation.json'))
    pending=[]
    for p in current_reports:
        if p.is_file() and 'PENDING_RTM_SYNCHRONIZATION' in p.read_text(encoding='utf-8'):
            pending.append(p.relative_to(ROOT).as_posix())
    add(checks,'current_machine_readable_validation_summaries_reflect_final_rtm_sync',not pending,{'pending_reports':pending,'reports_checked':len(current_reports)})
    errors += [f'pending-rtm:{x}' for x in pending]

    mapping_text=(ROOT/'conformance/test-mapping.json').read_text(encoding='utf-8')
    retired_active=[ns for ns in sorted(RETIRED_NS) if ns in mapping_text]
    add(checks,'retired_ns_not_active_test_mapping_targets',not retired_active,{'retired_ids_found':retired_active})
    if retired_active: errors.append('retired-ns-active-test-target')

    root_readme=(ROOT/'README.md').read_text(encoding='utf-8')
    inventory_ok=('393 permanent Normative Statement identifiers' in root_readme and '**390 are current**' in root_readme and 'RTM-1.47' in root_readme and 'RELEASE_NOTES_v2.1.x.md' in root_readme)
    add(checks,'root_readme_reflects_closed_v2_1_x_inventory',inventory_ok)
    if not inventory_ok: errors.append('readme-inventory')

    # Validate repository-relative Markdown links in files touched by closure.
    link_files=['README.md','PACKAGE_README.md','RELEASE_NOTES_v2.1.x.md','spec/README.md','conformance/README.md','governance/AGCP-v2.1.x-REPOSITORY-CLOSURE.md']
    link_re=re.compile(r'(?<!!)\[[^\]]*\]\(([^)]+)\)')
    link_issues=[]; links_checked=0
    for rel in link_files:
        p=ROOT/rel
        if not p.is_file(): continue
        for raw in link_re.findall(p.read_text(encoding='utf-8',errors='replace')):
            target=raw.strip().split()[0].strip('<>')
            if not target or target.startswith(('#','http://','https://','mailto:','tel:')): continue
            target=urllib.parse.unquote(target.split('#',1)[0].split('?',1)[0])
            if not target: continue
            links_checked+=1
            resolved=(p.parent/target).resolve()
            try: resolved.relative_to(ROOT)
            except ValueError:
                link_issues.append(f'outside:{rel}:{raw}'); continue
            if not resolved.exists(): link_issues.append(f'missing:{rel}:{raw}')
    add(checks,'closure_document_markdown_links_resolve',not link_issues,{'links_checked':links_checked,'issues':link_issues})
    errors += ['markdown:'+x for x in link_issues]

    # Repository closure manifest coverage and hashes.
    mf=j('governance/AGCP-v2.1.x-repository-closure-manifest.json')
    exclusions=set(mf.get('scope_exclusions',[]))
    entries={e['path']:e for e in mf.get('files',[])}
    actual=set()
    for p in ROOT.rglob('*'):
        if not p.is_file() or '.git' in p.parts: continue
        rel=p.relative_to(ROOT).as_posix()
        if rel in exclusions: continue
        actual.add(rel)
    missing_entries=sorted(actual-set(entries)); extra_entries=sorted(set(entries)-actual); mismatches=[]
    for rel,e in entries.items():
        p=ROOT/rel
        if not p.is_file() or sha256(p)!=e.get('sha256') or p.stat().st_size!=e.get('bytes'):
            mismatches.append(rel)
    manifest_ok=not (missing_entries or extra_entries or mismatches) and mf.get('file_count')==len(entries)
    add(checks,'repository_closure_manifest_complete_and_current',manifest_ok,{'manifest_entries':len(entries),'unlisted':missing_entries,'extra':extra_entries,'hash_or_size_mismatches':mismatches})
    if not manifest_ok: errors.append('repository-closure-manifest')

    status='PASS' if not errors else 'FAIL'
    report={
      'validation_id':'AGCP-v2.1.x-REPOSITORY-CLOSURE-VALIDATION-2026-10-03',
      'status':status,
      'validated_at':'2026-10-03',
      'synchronization_target':'v2.1.x',
      'controlling_published_baseline':'v2.1.0',
      'publication_version_assignment':'NOT_ASSIGNED_BY_CLOSURE_STEP',
      'rtm_dataset_version':'RTM-1.47',
      'checks_passed':sum(1 for c in checks if c['status']=='PASS'),
      'check_count':len(checks),
      'checks':checks,
      'warnings':warnings,
      'errors':errors,
      'next_controlled_step':'GENERIC_SOURCE_SYNCHRONIZATION' if status=='PASS' else 'RESOLVE_REPOSITORY_CLOSURE_ERRORS'
    }
    REPORT.write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
    print(json.dumps({'status':status,'checks_passed':report['checks_passed'],'check_count':len(checks),'errors':errors},indent=2))
    return 0 if status=='PASS' else 1

if __name__=='__main__': sys.exit(main())
